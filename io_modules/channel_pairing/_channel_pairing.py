"""Read and write the channel pairing workbook."""

import os

from openpyxl import Workbook, load_workbook

from selection.channel_identity import split_channel_base_and_direction
from selection.channel_pairing import resolve_channel_pair_rejection

PAIRING_FILE_NAME = "channel_pairing.xlsx"
_SHEET = "ChannelPairs"
_HEADER = ("Measured channel", "Simulated channel")


def _cell_name(identity) -> str:
    """A name that parses back to the same identity."""
    base, direction = identity
    return f"{base}:{direction}" if direction else base


def read_channel_pairing(path: str):
    """Return (pairs, warnings): valid pairs in row order, a message per skipped row.

    A missing file is no pairs and no warnings.
    """
    if not os.path.isfile(path):
        return [], []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[_SHEET] if _SHEET in workbook.sheetnames else workbook.worksheets[0]
            rows = list(sheet.iter_rows(min_row=1, max_col=2, values_only=True))
        finally:
            workbook.close()
    except Exception as error:  # a hand-broken workbook must not block opening the project
        return [], [f"{os.path.basename(path)} cannot be read ({error})"]
    pairs, warnings = [], []
    for number, row in enumerate(rows, start=1):
        # A hand-made sheet may have no header: only a header-looking first row is skipped.
        if number == 1 and str(row[0] or "").strip().lower() == _HEADER[0].lower():
            continue
        names = [str(cell).strip() if cell is not None else "" for cell in row]
        names += [""] * (2 - len(names))
        if not any(names):
            continue
        if not all(names):
            warnings.append(f"row {number} - a channel is missing")
            continue
        pair = tuple(split_channel_base_and_direction(name) for name in names)
        reason = resolve_channel_pair_rejection(pair, pairs)
        if reason:
            warnings.append(f"row {number} - {reason}")
        else:
            pairs.append(pair)
    return pairs, warnings


def write_channel_pairing(path: str, pairs) -> None:
    """Replace the workbook with `pairs`; a locked file raises and stays intact."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _SHEET
    sheet.append(_HEADER)
    for left, right in pairs:
        sheet.append((_cell_name(left), _cell_name(right)))
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    temp_path = path + ".tmp"
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    except Exception:
        if os.path.exists(temp_path):
            os.remove(temp_path)
        raise
