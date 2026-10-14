"""
MASTA result (.xlsx) reader: order amplitude against speed per housing
location and axis, read as an Imported result (§1.134).

A workbook with sheets named "H<order>" (H1, H4.79, H22, ...) holds one result
per sheet; every such sheet is read and any other sheet (ERP) is ignored. A
workbook without them is read from its first sheet. Each channel remembers its
sheet, and its index is the sheet's place among the read sheets times
_SHEET_STRIDE plus its column pair, so the first sheet keeps the indices it had
when only that sheet was read (§1.159).

Layout of each read sheet -- every channel is a pair of columns, each with its
own speed column, because MASTA's speed steps are not shared between pairs:

    row 1  "<axis>, At housing: <design>\\<location>"     -> name "<location>:<axis>"
    row 2  result type, damping included, "... Order <n>-Damping = ..."
    row 3  scenario
    row 4  "Speed (rev/min)" | "Amplitude (<unit>)"
    row 5+ rpm | amplitude, either of them possibly stored as text

The header is Raw metadata of each channel, read at scan: design, scenario,
damping and description_1..N (row 2's comma segments left once the Order and
Damping keywords are taken out). The order goes to the block, not a facet.
"""

import functools
import os
import re
from typing import List, Optional, Tuple

import numpy as np

from core.block_kinds import KIND_ORDER_CUT, STEP_READ
from core.data_block import Acquisition, NVHDataBlock, Provenance
from core.models import FUNC_TYPE_ORDER_FUNCTION, ChannelMetadata, MeasurementRunIndex
from core.units import determine_channel_type
from io_modules.measurement_files import canonical_path, file_stamp

_HEADER_ROWS = 4
_LOCATION_RE = re.compile(r"^\s*([XYZ])\s*,.*\\([^\\]+?)\s*$")
_DESIGN_RE = re.compile(r"At housing:\s*(.*)\\")
_SPEED_HEADER = "speed (rev/min)"
_UNIT_RE = re.compile(r"\(([^)]*)\)\s*$")
_ORDER_RE = re.compile(r"\bOrder\s+(\d+(?:\.\d+)?)")
_DAMPING_RE = re.compile(r"\bDamping\s*=\s*([^\s,]+)")
_SEGMENT_EDGES = " \t-–:;"
_ORDER_SHEET_RE = re.compile(r"^H\d+(?:\.\d+)?$", re.IGNORECASE)
# ponytail: assumes fewer than 1000 column pairs per sheet; widen if MASTA ever exports more.
_SHEET_STRIDE = 1000


def _parse_result_type(text: str) -> Tuple[Optional[float], str, List[str]]:
    """Row 2 as (order, damping, descriptions); a missing keyword is None / ""."""
    order_match = _ORDER_RE.search(text)
    damping_match = _DAMPING_RE.search(text)
    descriptions = []
    for segment in text.split(","):
        segment = _DAMPING_RE.sub("", _ORDER_RE.sub("", segment)).strip(_SEGMENT_EDGES)
        if segment:
            descriptions.append(segment)
    return (float(order_match.group(1)) if order_match else None,
            damping_match.group(1) if damping_match else "",
            descriptions)


def _read_sheets(file_path: str, max_row: Optional[int] = None,
                 only: Optional[int] = None) -> List[Tuple[Optional[str], List[tuple]]]:
    """(sheet name, rows) of every read sheet, or of the `only`-th; name None for a first-sheet workbook."""
    import openpyxl
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        named = [(sheet.title.strip(), sheet) for sheet in workbook.worksheets
                 if _ORDER_SHEET_RE.match(sheet.title.strip())]
        sheets = named or [(None, workbook.worksheets[0])]
        if only is not None:
            sheets = sheets[only:only + 1]
        return [(name, list(sheet.iter_rows(max_row=max_row, values_only=True)))
                for name, sheet in sheets]
    finally:
        workbook.close()


@functools.lru_cache(maxsize=2)
def _read_all_sheets(path: str, _mtime_ns: int, _size: int) -> List[Tuple[Optional[str], List[tuple]]]:
    """
    Every read sheet, parsed once per file stamp: DataAccessor opens a fresh
    reader per channel, and openpyxl's load_workbook walks every sheet each time,
    so a 22-channel plot paid for 22 full loads.
    """
    return _read_sheets(path)


def _read_sheets_cached(file_path: str) -> List[Tuple[Optional[str], List[tuple]]]:
    stamp = file_stamp(file_path)
    if stamp is None:
        return _read_sheets(file_path)
    return _read_all_sheets(canonical_path(file_path), stamp["mtime_ns"], stamp["size"])


def discard_masta_caches() -> None:
    """Forget the parsed workbook, so the next read parses it again (a Benchmark case starts cold)."""
    _read_all_sheets.cache_clear()


def is_masta_file(file_path: str) -> bool:
    """Whether the workbook's header is a MASTA result's. Any failure to read it is a no."""
    try:
        rows = _read_sheets(file_path, max_row=_HEADER_ROWS, only=0)[0][1]
    except Exception:
        return False
    return (len(rows) == _HEADER_ROWS and bool(rows[0]) and bool(rows[3])
            and _LOCATION_RE.match(str(rows[0][0] or "")) is not None
            and str(rows[3][0] or "").strip().lower() == _SPEED_HEADER)


def _parse_number(cell) -> Optional[float]:
    """A cell as a number, or None for a blank or a non-number (a footer, a note)."""
    if isinstance(cell, str):
        cell = cell.strip()
        if "," in cell and "." not in cell:
            cell = cell.replace(",", ".")   # Number stored as text with a decimal comma.
    try:
        return float(cell)
    except (TypeError, ValueError):
        return None


def _resolve_channel_indices(rows) -> List[int]:
    """Column pairs headed by a housing location; a stray trailing column is not a channel."""
    return [index for index in range(len(rows[0]) // 2)
            if _LOCATION_RE.match(str(rows[0][2 * index] or "")) is not None]


class AccReaderMasta:
    """Parser for MASTA result workbooks; one channel per (location, axis) column pair."""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def scan_file_structure(self) -> MeasurementRunIndex:
        return MeasurementRunIndex(os.path.basename(self.file_path), self.file_path)

    def load_channels_on_demand(self, run_index: MeasurementRunIndex):
        if run_index.available_channels:
            return
        for position, (sheet, rows) in enumerate(_read_sheets_cached(self.file_path)):
            for pair in _resolve_channel_indices(rows):
                index = position * _SHEET_STRIDE + pair
                name, unit = self._resolve_name_and_unit(rows, pair)
                meta = ChannelMetadata(index=index, name=name,
                                       channel_type=determine_channel_type(unit, name), unit=unit)
                meta.func_type = FUNC_TYPE_ORDER_FUNCTION
                meta.block_kind = KIND_ORDER_CUT
                meta.num_pts = len(self._read_pairs(rows, pair))
                design_match = _DESIGN_RE.search(str(rows[0][2 * pair]))
                meta.design = design_match.group(1).strip() if design_match else ""
                meta.scenario = str(rows[2][2 * pair] or "").strip()
                _, meta.damping, descriptions = _parse_result_type(str(rows[1][2 * pair] or ""))
                for number, description in enumerate(descriptions, start=1):
                    setattr(meta, f"description_{number}", description)
                if sheet is None:
                    run_index.add_channel_metadata(f"Col #{index}: {name}", meta)
                else:
                    meta.sheet = sheet
                    run_index.add_channel_metadata(f"Col #{index} [{sheet}]: {name}", meta)

    @staticmethod
    def _resolve_name_and_unit(rows, index: int):
        axis, location = _LOCATION_RE.match(str(rows[0][2 * index])).groups()
        unit_match = _UNIT_RE.search(str(rows[3][2 * index + 1] or ""))
        return f"{location}:{axis}", (unit_match.group(1).strip() if unit_match else "unknown")

    @staticmethod
    def _read_pairs(rows, index: int) -> List[tuple]:
        """(rpm, amplitude) of one channel; a pair ends where its columns run out."""
        pairs = []
        for row in rows[_HEADER_ROWS:]:
            cells = row[2 * index: 2 * index + 2]
            if len(cells) < 2:
                continue
            rpm, amplitude = _parse_number(cells[0]), _parse_number(cells[1])
            if rpm is not None and amplitude is not None:
                pairs.append((rpm, amplitude))
        return pairs

    def read_single_channel_data(self, dataset_index: int) -> NVHDataBlock:
        """
        A finished order_cut block, not a RawChannel: a RawChannel is a
        waveform. DataAccessor still adds where it came from and its channel type.
        """
        position, pair = divmod(dataset_index, _SHEET_STRIDE)
        sheets = _read_sheets_cached(self.file_path)[position:position + 1]
        rows = sheets[0][1] if sheets else []
        if not rows or pair not in _resolve_channel_indices(rows):
            raise IndexError(
                f"Channel index {dataset_index} is not a channel of {os.path.basename(self.file_path)}.")
        name, unit = self._resolve_name_and_unit(rows, pair)
        pairs = np.asarray(self._read_pairs(rows, pair), dtype=np.float64).reshape(-1, 2)
        order, _, _ = _parse_result_type(str(rows[1][2 * pair] or ""))
        return NVHDataBlock.order_cut(
            name=name, values=pairs[:, 1], rpm=pairs[:, 0], order=order,
            value_unit=unit, processing=None,
            acquisition=Acquisition(num_pts=len(pairs), recorded_unit=unit),
            provenance=Provenance(step=STEP_READ),
        )
